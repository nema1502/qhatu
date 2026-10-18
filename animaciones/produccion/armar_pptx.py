"""Arma deck/qhatu-deck.pptx para exponer en vivo.

Una lámina por diapositiva, con su animación incrustada (se reproduce al hacer clic;
mientras tanto se ve la lámina terminada) y, en las notas del orador, cada frase del
guion con quién la lee. La diapositiva de validación lleva las respuestas cortas para
las preguntas del jurado.

    python animaciones/producir_pitch.py pptx
"""

import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent.parent
sys.path.insert(0, str(AQUI.parent))

import guion  # noqa: E402

LAMINAS = [("S1", "S1Portada", "01-portada"), ("S2", "S2Problema", "02-problema"),
           ("S3", "S3Solucion", "03-solucion"), ("S4", "S4Flujo", "04-flujo"),
           ("S5", "S5Stellar", "05-stellar"), ("S6", "S6Modelo", "06-modelo"),
           ("S7", "S7Validacion", "07-validacion"), ("S8", "S8Plan", "08-plan")]

NOTAS_S7 = """Para las preguntas del jurado. Las respuestas completas están en docs/defensa.md.

- ¿Es lo mismo que Trustless Work? No: ellos son la caja fuerte (escrow no custodial); Qhatu es el producto de comercio exterior encima.
- ¿Más barato que una carta de crédito? No lo afirmamos: del mismo orden (0,75% a 3% contra 1% a 1,5%). Buscamos acceso y rapidez, y lo medimos en el piloto.
- ¿Por qué el comprador deposita el 100%? Es la mayor duda del diseño; por eso se libera en dos hitos.
- ¿Es legal? Falta opinión legal. Qhatu no toca los fondos y habla de pago condicionado por hitos.
- ¿Entrevistas? No formales todavía: la evidencia es pública. Siguen entrevistas y 3 tratos reales.
- ¿Qué hay hecho? El diseño, el deck, el video y una maqueta. La demo en testnet es lo próximo."""


def clip(trabajo, escena):
    base = trabajo / ("s4" if escena == "S4Flujo" else f"manim/{escena}")
    return next(base.rglob(f"{escena}.mp4"), None)


def main(trabajo):
    trabajo = Path(trabajo)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    en_video = {esc: guion.ESCENAS[esc] for esc in guion.ORDEN}
    for esc, escena, png in LAMINAS:
        imagen = RAIZ / "deck" / "laminas" / f"{png}.png"
        diapo = prs.slides.add_slide(prs.slide_layouts[6])
        diapo.shapes.add_picture(str(imagen), 0, 0, width=prs.slide_width, height=prs.slide_height)
        video = clip(trabajo, escena)
        if video:
            diapo.shapes.add_movie(str(video), 0, 0, prs.slide_width, prs.slide_height,
                                   poster_frame_image=str(imagen), mime_type="video/mp4")
        if esc in en_video:
            notas = "\n\n".join(f"{guion.quien(esc, i)}: {t}" for i, t, _ in en_video[esc])
        else:
            notas = NOTAS_S7
        diapo.notes_slide.notes_text_frame.text = notas
    destino = RAIZ / "deck" / "qhatu-deck.pptx"
    prs.save(destino)
    print(f"{destino} ({destino.stat().st_size / 1048576:.1f} MB)")


if __name__ == "__main__":
    main(sys.argv[1])
